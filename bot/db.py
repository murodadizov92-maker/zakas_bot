"""Android ilova uchun yangi API (/api/v2/...). Eski Telegram Mini App yo'llari (/api/products, /api/order) o'zgarmaydi."""
import hashlib
import hmac
import html
import io
import json
import logging
import os
import re
import secrets
import time
from datetime import datetime, timedelta

import aiohttp
from aiohttp import web

import db
from verify import verify_init_data
from config import (
    SECRET_KEY, TIMEZONE, ADMIN_CHAT_ID, ADMIN_PHONE,
    SUPABASE_URL, SUPABASE_KEY, SUPABASE_BUCKET,
)

log = logging.getLogger("zakas-api")

# Avtomobillar: max_kg — bitta zakas hajmi chegarasi, min_kg — kamida shuncha bo'lishi shart,
# weekly — bitta filial haftada faqat 1 marta (admin qo'shimcha bersa bundan tashqari ham mumkin)
VEHICLES = {
    "labo": {"name": "Labo", "max_kg": 2000, "min_kg": 0, "weekly": False, "shared": False},   # 1000 kg × 2 tagacha
    "shinaray": {"name": "Shinaray", "max_kg": 2200, "min_kg": 0, "weekly": True, "shared": True},
    "isuzi": {"name": "Isuzi", "max_kg": 10000, "min_kg": 0, "weekly": False, "shared": True},
}
LABO_CAP = 1000   # bitta Labo sig'imi, kg
MAX_KG = 10000    # bitta zakasning eng katta hajmi


def fits(veh: str, n: int, eq: float) -> bool:
    """Admin tasdiqlashida: yuk shu avtomobil(lar)ga sig'adimi? (faqat sig'im tekshiriladi, qolganini admin hal qiladi)"""
    if veh == "labo":
        return 1 <= n <= 10 and eq <= LABO_CAP * n + 1e-6
    if veh == "shinaray":
        return n == 1 and eq <= VEHICLES["shinaray"]["max_kg"] + 1e-6
    if veh == "isuzi":
        return n == 1 and eq <= MAX_KG + 1e-6
    return False


def vehicle_ok(veh: str, n: int, eq: float) -> bool:
    """Zakas hajmiga mos avtomobilmi? (ilovadagi suggest() bilan bir xil qoida)
    ≤1000 — 1 Labo; ≤2000 — 2 Labo; 2000–2200 — Shinaray yoki Isuzi; 2200–10000 — Isuzi."""
    if veh == "labo":
        return n in (1, 2) and eq <= LABO_CAP * n + 1e-6 and (n == 1 or eq > LABO_CAP)
    if veh == "shinaray":
        return n == 1 and 2000 < eq <= 2200 + 1e-6
    if veh == "isuzi":
        return n == 1 and 2000 < eq <= MAX_KG + 1e-6
    return False


DONA_KG = 0.4  # 1 dona = 0.4 kg


def quota_error(day_used: bool, week_used: bool, taken: bool, has_grant: bool):
    """Oddiy huquq yetmasa, admin bergan qo'shimcha huquq (grant) kerak. Xato kodi yoki None."""
    if not (day_used or week_used or taken) or has_grant:
        return None
    if taken:
        return "vehicle_taken"
    return "daily_limit" if day_used else "weekly_limit"


def eq_kg(items) -> float:
    """Zakasning kg dagi hajmi: kg lar o'z holicha, donalar × 0.4."""
    return round(sum(i["qty"] * (DONA_KG if i.get("unit") == "dona" else 1) for i in items), 2)

BOT = None  # main.py o'rnatadi


def chat_ids() -> list:
    """Zakaslar boradigan chatlar: ADMIN_CHAT_ID + (ixtiyoriy) GROUP_CHAT_ID — vergul bilan bir nechta."""
    out = []
    for x in [ADMIN_CHAT_ID] + os.getenv("GROUP_CHAT_ID", "").split(","):
        try:
            v = int(str(x).strip())
        except (TypeError, ValueError):
            continue
        if v and v not in out:
            out.append(v)
    return out


def chunks(text: str, n: int = 3800) -> list:
    out, cur = [], ""
    for line in text.split("\n"):
        if cur and len(cur) + len(line) + 1 > n:
            out.append(cur.rstrip())
            cur = ""
        cur += line + "\n"
    if cur.strip():
        out.append(cur.rstrip())
    return out


async def notify(text: str) -> int:
    """Matnni chatlarga yuboradi. Muvaffaqiyatli yuborilgan xabarlar soni qaytadi."""
    sent = 0
    if not BOT:
        return sent
    for cid in chat_ids():
        for part in chunks(text):
            try:
                await BOT.send_message(cid, part, parse_mode="HTML")
                sent += 1
            except Exception as e:
                log.warning("Telegramga (%s) yuborib bo'lmadi: %s", cid, e)
    return sent


def short_name(name: str) -> str:
    return re.sub(r"\s*\*[^*]*\*\s*$", "", str(name)).strip()


def order_text(oid, d, vname, vmax, u, items, kg, dona, eq) -> str:
    """Zakas matni: mahsulotlar kategoriya bo'yicha ketma-ket taxlangan."""
    head = (
        f"🆕 <b>Yangi buyurtma #{oid}</b>\n"
        f"📅 Sana: {d.strftime('%d.%m.%Y')}\n"
        f"🚚 Avto: <b>{vname}</b>\n"
        f"👤 {html.escape(u['name'])} (+998{u['phone']})\n"
    )
    body, cur = [], None
    for it in items:  # items allaqachon kategoriya/tartib bo'yicha saralangan
        if it["cat"] != cur:
            cur = it["cat"]
            body.append(f"\n📂 <b>{html.escape(cur)}</b>")
        body.append(f"• {html.escape(short_name(it['name']))} — <b>{it['qty']:g} {it['unit']}</b>")
    tail = f"\n\n<b>Jami:</b> {kg:g} kg, {dona:g} dona (≈ {eq:g} / {vmax} kg)"
    return head + "\n".join(body) + tail
TOKEN_TTL = 30 * 24 * 3600
FAILS: dict[str, list[float]] = {}  # telefon -> xato urinishlar vaqti


# ------------------------------------------------------------------ yordamchilar
def make_token(uid: int) -> str:
    exp = int(time.time()) + TOKEN_TTL
    msg = f"{uid}.{exp}"
    sig = hmac.new(SECRET_KEY.encode(), msg.encode(), hashlib.sha256).hexdigest()
    return f"{msg}.{sig}"


def parse_token(token: str) -> int | None:
    try:
        uid, exp, sig = token.split(".")
        msg = f"{uid}.{exp}"
        good = hmac.new(SECRET_KEY.encode(), msg.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(good, sig) or int(exp) < time.time():
            return None
        return int(uid)
    except Exception:
        return None


def err(code: str, status: int = 400) -> web.Response:
    return web.json_response({"ok": False, "error": code}, status=status)


def user_json(u, with_code: bool = False) -> dict:
    d = {
        "id": u["id"], "name": u["name"], "phone": "+998" + u["phone"],
        "is_admin": u["is_admin"], "active": u["active"],
    }
    if with_code:
        d["code"] = u["code"]
    return d


def tg_user_id(init_data: str) -> int | None:
    """Telegram initData to'g'ri va yangi bo'lsa — foydalanuvchi ID'si, aks holda None."""
    pairs = verify_init_data(init_data or "")
    if pairs is None:
        return None
    try:
        if time.time() - int(pairs.get("auth_date", 0)) > 7 * 24 * 3600:
            return None
        return int(json.loads(pairs.get("user", "{}")).get("id", 0)) or None
    except (ValueError, TypeError):
        return None


def hhmm_to_min(s: str) -> int:
    h, m = s.split(":")
    return int(h) * 60 + int(m)


async def order_status() -> dict:
    s = await db.get_settings()
    now = datetime.now(TIMEZONE)
    cur = now.hour * 60 + now.minute
    is_open = hhmm_to_min(s["open_time"]) <= cur < hhmm_to_min(s["cutoff_time"])
    days = None  # None — admin hali kunlarni belgilamagan (Yakshanbadan boshqa hammasi ochiq)
    if "order_days" in s:
        try:
            days = sorted({str(x) for x in json.loads(s["order_days"])})
        except (ValueError, TypeError):
            days = None
    return {"open_time": s["open_time"], "cutoff_time": s["cutoff_time"], "order_open": is_open,
            "order_days": days}


def norm_name(s: str) -> str:
    return re.sub(r"[\W_]+", "", str(s).lower())


# ------------------------------------------------------------------ middleware
CORS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Headers": "Authorization, Content-Type",
    "Access-Control-Allow-Methods": "GET, POST, PUT, DELETE, OPTIONS",
}


@web.middleware
async def cors_mw(request, handler):
    if not request.path.startswith("/api/v2"):
        return await handler(request)
    if request.method == "OPTIONS":
        return web.Response(status=204, headers=CORS)
    resp = await handler(request)
    resp.headers.update(CORS)
    return resp


PUBLIC = {"/api/v2/check", "/api/v2/login", "/api/v2/tg-login"}


@web.middleware
async def auth_mw(request, handler):
    path = request.path
    if not path.startswith("/api/v2") or path in PUBLIC:
        return await handler(request)
    auth = request.headers.get("Authorization", "")
    uid = parse_token(auth.removeprefix("Bearer ").strip())
    if uid is None:
        return err("unauthorized", 401)
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where id=$1", uid)
    if not u or not u["active"]:
        return err("not_allowed", 403)
    if path.startswith("/api/v2/admin") and not u["is_admin"]:
        return err("forbidden", 403)
    request["user"] = u
    return await handler(request)


# ------------------------------------------------------------------ kirish
async def h_check(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where phone=$1", phone)
    if not u or not u["active"]:
        return web.json_response({"ok": True, "status": "denied"})
    return web.json_response({"ok": True, "status": "new" if not u["code"] else "has_code"})


async def h_login(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    code = str(data.get("code") or "").strip()

    now = time.time()
    FAILS[phone] = [t for t in FAILS.get(phone, []) if now - t < 600]
    if len(FAILS[phone]) >= 5:
        return err("too_many_attempts", 429)

    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where phone=$1", phone)
        if not u or not u["active"]:
            return err("not_allowed", 403)
        if not u["code"]:  # birinchi kirish: foydalanuvchi o'z kodini yozadi
            if len(code) < 4:
                return err("code_too_short", 400)
            await c.execute("update users set code=$1 where id=$2", code, u["id"])
            u = await c.fetchrow("select * from users where id=$1", u["id"])
        elif not hmac.compare_digest(u["code"], code):
            FAILS[phone].append(now)
            return err("wrong_code", 401)
        tg = tg_user_id(str(data.get("initData") or ""))
        if tg:  # Telegram ichidan kirdi: keyingi safar kod so'ramasligi uchun bog'laymiz
            await c.execute("update users set tg_id=null where tg_id=$1", tg)
            await c.execute("update users set tg_id=$1 where id=$2", tg, u["id"])
    FAILS.pop(phone, None)
    return web.json_response({"ok": True, "token": make_token(u["id"]), "user": user_json(u)})


async def h_tg_login(request):
    """Telegram Mini App: bog'langan foydalanuvchi (yoki ADMIN_CHAT_ID egasi) kodsiz kiradi."""
    data = await request.json()
    tg_id = tg_user_id(str(data.get("initData") or ""))
    if tg_id is None:
        return err("auth_failed", 401)
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where tg_id=$1 and active", tg_id)
        if not u and ADMIN_CHAT_ID and tg_id == ADMIN_CHAT_ID:
            u = await c.fetchrow(
                "select * from users where phone=$1 and is_admin and active", db.norm_phone(ADMIN_PHONE))
    if not u:
        return err("not_allowed", 403)
    return web.json_response({"ok": True, "token": make_token(u["id"]), "user": user_json(u)})


async def h_me(request):
    return web.json_response({"ok": True, "user": user_json(request["user"])})


# ------------------------------------------------------------------ katalog va buyurtma
async def h_products(request):
    extra = {}
    async with db.pool.acquire() as c:
        rows = await c.fetch(
            "select id,name,price,category,image_url from products where active order by sort,id"
        )
        if request["user"]["is_admin"]:  # tasdiq kutayotgan zakaslar soni (Sozlamalar belgisi uchun)
            extra["pending"] = await c.fetchval("select count(*) from orders where status='pending'")
    catalog: dict[str, list] = {}
    for r in rows:
        catalog.setdefault(r["category"], []).append(
            {"id": r["id"], "name": r["name"], "price": r["price"], "image": r["image_url"]}
        )
    return web.json_response({"ok": True, "catalog": catalog, **extra, **(await order_status())})


async def h_order(request):
    u = request["user"]
    data = await request.json()
    st = await order_status()
    if not st["order_open"]:
        return err("closed", 403)

    veh = data.get("vehicle")
    if veh not in VEHICLES:
        return err("bad_vehicle")
    V = VEHICLES[veh]
    try:
        veh_n = int(data.get("vehicle_n") or 1)
    except (TypeError, ValueError):
        return err("bad_vehicle")

    today = datetime.now(TIMEZONE).date()
    try:
        d = datetime.strptime(str(data.get("date") or today.isoformat()), "%Y-%m-%d").date()
    except ValueError:
        return err("bad_date")
    if d < today or d > today + timedelta(days=14):
        return err("bad_date")
    if st["order_days"] is not None and d.isoformat() not in st["order_days"]:
        return err("day_closed")

    async with db.pool.acquire() as c:
        prods = {r["id"]: r for r in await c.fetch("select id,name,price,category,sort from products where active")}
    lines, items, kg, dona = [], [], 0.0, 0.0
    for it in data.get("items") or []:
        p = prods.get(it.get("id"))
        try:
            qty = float(it.get("qty", 0))
        except (TypeError, ValueError):
            continue
        if not p or qty <= 0:
            continue
        unit = "dona" if it.get("unit") == "dona" else "kg"
        kg += qty if unit == "kg" else 0
        dona += qty if unit == "dona" else 0
        items.append({"id": p["id"], "name": p["name"], "price": p["price"], "qty": qty, "unit": unit,
                      "cat": p["category"]})
    if not items:
        return err("empty_order")
    items.sort(key=lambda x: (prods[x["id"]]["sort"], x["id"]))  # kategoriya va katalog tartibida

    eq = eq_kg(items)
    if eq > MAX_KG + 1e-6:
        return err("limit_full")
    if not vehicle_ok(veh, veh_n, eq):  # tanlangan avto zakas hajmiga mos emas
        return err("vehicle_mismatch")

    # Foydalanuvchi cheklovsiz zakas bera oladi: avtoni admin «Buyurtmalar» bo'limida tanlab tasdiqlaydi.
    # Zakas guruhga FAQAT tasdiqlangandan keyin boradi (h_order_confirm). Tanlangan avto — shunchaki so'rov.
    async with db.pool.acquire() as c:
        oid = await c.fetchval(
            "insert into orders(user_id,order_date,items,total_kg,total_dona,vehicle,vehicle_n,status,asked_vehicle,asked_n) "
            "values($1,$2,$3::jsonb,$4,$5,$6,$7,'pending',$6,$7) returning id",
            u["id"], d, json.dumps(items, ensure_ascii=False), kg, dona, veh, veh_n,
        )
    return web.json_response({"ok": True, "order_id": oid, "status": "pending"})


async def vehicle_availability(c, u, d) -> dict:
    """Har avto uchun shu foydalanuvchi shu kunga zakas bera oladimi: {veh: {"ok": bool, "reason": kod|None}}."""
    gr = {r["vehicle"]: r["n"] for r in await c.fetch(
        "select vehicle, count(*) as n from grants where user_id=$1 and used_order_id is null group by vehicle", u["id"])}
    day_used = (await c.fetchval(
        "select count(*) from orders where user_id=$1 and order_date=$2 and grant_id is null", u["id"], d)) > 0
    out = {}
    for k, V in VEHICLES.items():
        week_used = False
        if V["weekly"]:
            ws = d - timedelta(days=d.weekday())
            week_used = (await c.fetchval(
                "select count(*) from orders where user_id=$1 and vehicle=$2 and grant_id is null "
                "and order_date>=$3 and order_date<$4", u["id"], k, ws, ws + timedelta(days=7))) > 0
        taken = False
        if V["shared"]:
            taken = (await c.fetchval(
                "select count(*) from orders where order_date=$1 and vehicle=$2 and user_id<>$3",
                d, k, u["id"])) > 0
        code = quota_error(day_used, week_used, taken, gr.get(k, 0) > 0)
        out[k] = {"ok": code is None, "reason": code}
    return out


async def h_vehicles(request):
    """Avtomobillar sozlamasi, qo'shimcha huquqlar va (date berilsa) shu kunga har avtoning holati."""
    d = None
    if request.query.get("date"):
        try:
            d = datetime.strptime(request.query["date"], "%Y-%m-%d").date()
        except ValueError:
            return err("bad_date")
    async with db.pool.acquire() as c:
        rows = await c.fetch(
            "select vehicle, count(*) as n from grants where user_id=$1 and used_order_id is null group by vehicle",
            request["user"]["id"])
        avail = await vehicle_availability(c, request["user"], d) if d else {}
    return web.json_response({
        "ok": True,
        "vehicles": {k: {"name": v["name"], "max_kg": v["max_kg"], "min_kg": v["min_kg"], "weekly": v["weekly"]}
                     for k, v in VEHICLES.items()},
        "grants": {r["vehicle"]: r["n"] for r in rows},
        "avail": avail,
    })


async def h_request(request):
    """«＋» tugmasi: foydalanuvchi xabar yozadi — adminga o'sha foydalanuvchi nomidan zakas kabi boradi."""
    u = request["user"]
    data = await request.json()
    text = str(data.get("text") or "").strip()[:1000]
    if not text:
        return err("empty_text")
    async with db.pool.acquire() as c:
        rid = await c.fetchval("insert into requests(user_id,text) values($1,$2) returning id", u["id"], text)
    await notify(
        f"➕ <b>Qo'shimcha avto so'rovi #{rid}</b>\n"
        f"📅 Sana: {datetime.now(TIMEZONE).strftime('%d.%m.%Y %H:%M')}\n"
        f"👤 {html.escape(u['name'])} (+998{u['phone']})\n\n"
        f"<b>Xabar:</b>\n{html.escape(text)}"
    )
    return web.json_response({"ok": True, "id": rid})


async def h_my_orders(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch(
            "select id,order_date,created_at,items,total_kg,total_dona,vehicle,vehicle_n,status,asked_vehicle,asked_n "
            "from orders where user_id=$1 order by id desc limit 50", request["user"]["id"],
        )
    return web.json_response({"ok": True, "orders": [order_json(r) for r in rows]})


def order_json(r) -> dict:
    return {
        "id": r["id"], "date": r["order_date"].isoformat(),
        "created_at": r["created_at"].astimezone(TIMEZONE).strftime("%Y-%m-%d %H:%M"),
        "items": json.loads(r["items"]) if isinstance(r["items"], str) else r["items"],
        "total_kg": float(r["total_kg"]), "total_dona": float(r["total_dona"]),
        "vehicle": r["vehicle"], "vehicle_n": r["vehicle_n"],
        "status": r["status"], "asked_vehicle": r["asked_vehicle"], "asked_n": r["asked_n"],
    }


# ------------------------------------------------------------------ ADMIN: foydalanuvchilar
async def h_users(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch("select * from users order by id")
    return web.json_response({"ok": True, "users": [user_json(r, True) for r in rows]})


async def h_user_add(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    name = str(data.get("name") or "").strip()
    if len(phone) != 9 or not name:
        return err("bad_input")
    async with db.pool.acquire() as c:
        if await c.fetchval("select 1 from users where phone=$1", phone):
            return err("exists", 409)
        uid = await c.fetchval(
            "insert into users(phone,name) values($1,$2) returning id", phone, name
        )
    return web.json_response({"ok": True, "id": uid})


async def h_user_update(request):
    uid = int(request.match_info["id"])
    data = await request.json()
    sets, vals = [], []

    def add(col, val):
        vals.append(val)
        sets.append(f"{col}=${len(vals)}")

    if "name" in data:
        add("name", str(data["name"]).strip())
    if "active" in data:
        if uid == request["user"]["id"] and not data["active"]:
            return err("cannot_disable_self")
        add("active", bool(data["active"]))
    if "is_admin" in data:
        if uid == request["user"]["id"] and not data["is_admin"]:
            return err("cannot_demote_self")
        add("is_admin", bool(data["is_admin"]))
    if data.get("generate_code") or "code" in data:
        add("tg_id", None)  # kod o'zgarsa Telegram'dan qayta kod bilan kirishi kerak
    if data.get("generate_code"):
        add("code", f"{secrets.randbelow(10000):04d}")
    elif "code" in data:  # None yoki "" => kodni tozalash (foydalanuvchi o'zi qayta yozadi)
        c_ = str(data["code"] or "").strip()
        if c_ and len(c_) < 4:
            return err("code_too_short")
        add("code", c_ or None)
    if not sets:
        return err("nothing_to_update")
    vals.append(uid)
    async with db.pool.acquire() as c:
        u = await c.fetchrow(
            f"update users set {', '.join(sets)} where id=${len(vals)} returning *", *vals
        )
    if not u:
        return err("not_found", 404)
    return web.json_response({"ok": True, "user": user_json(u, True)})


# ------------------------------------------------------------------ ADMIN: mahsulot, narx
async def h_admin_products(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch("select * from products order by sort,id")
    return web.json_response({"ok": True, "products": [
        {"id": r["id"], "name": r["name"], "price": r["price"], "category": r["category"],
         "image": r["image_url"], "active": r["active"]} for r in rows]})


async def h_product_update(request):
    pid = int(request.match_info["id"])
    data = await request.json()
    sets, vals = [], []
    for col, cast in (("price", int), ("name", str), ("category", str), ("active", bool)):
        if col in data:
            vals.append(cast(data[col]))
            sets.append(f"{col}=${len(vals)}")
    if not sets:
        return err("nothing_to_update")
    vals.append(pid)
    async with db.pool.acquire() as c:
        r = await c.fetchrow(
            f"update products set {', '.join(sets)} where id=${len(vals)} returning id", *vals
        )
    return web.json_response({"ok": True}) if r else err("not_found", 404)


_HOMO = str.maketrans("АВЕКМНОРСТХаеорсх", "ABEKMHOPCTXaeopcx")  # kirill → o'xshash lotin harflar (ЕZOZ = EZOZ)


def cat_key(s) -> str:
    return re.sub(r"[\W_]+", "", str(s).translate(_HOMO).lower())


def clean_cat(s) -> str:
    return re.sub(r"\s+", " ", str(s).replace("*", " ")).strip()


def parse_price_excel(raw: bytes) -> list:
    """Excel'dan (nom, narx|None, kategoriya|None) qatorlarini oladi. Sarlavha qatorini o'zi topadi.
    Narxsiz qatorlar ham olinadi (kategoriyani yangilash uchun), narxi esa None bo'ladi."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    name_keys = ("наимен", "номи", "nom", "name", "товар", "mahsulot", "product")
    price_keys = ("цена", "прайс", "narx", "price", "нарх", "стоим")
    cat_keys = ("категор", "kategor", "categor", "группа", "guruh")
    name_col = price_col = cat_col = hdr = None
    for i, row in enumerate(rows[:15]):
        n = p = c = None
        for j, cell in enumerate(row):
            t = str(cell or "").lower()
            if n is None and any(k in t for k in name_keys):
                n = j
            elif c is None and any(k in t for k in cat_keys):
                c = j
            elif p is None and any(k in t for k in price_keys):
                p = j
        if n is not None and p is not None:
            name_col, price_col, cat_col, hdr = n, p, c, i
            break
    if hdr is None:  # zaxira: birinchi matn ustuni va birinchi son ustuni
        hdr = 0
        sample = rows[1] if len(rows) > 1 else []
        name_col = next((j for j, v in enumerate(sample) if isinstance(v, str)), 0)
        price_col = next((j for j, v in enumerate(sample) if isinstance(v, (int, float))), 1)
    out = []
    for row in rows[hdr + 1:]:
        if len(row) <= max(name_col, price_col):
            continue
        name = row[name_col]
        if not name:
            continue
        try:
            price = int(round(float(str(row[price_col]).replace(" ", "").replace(",", "."))))
        except (TypeError, ValueError):
            price = None
        if price is not None and price <= 0:
            price = None
        cat = None
        if cat_col is not None and len(row) > cat_col and row[cat_col]:
            cat = clean_cat(row[cat_col]) or None
        out.append((str(name).strip(), price, cat))
    return out


def plan_import(rows: list, existing: list, add_new: bool) -> dict:
    """Import rejasi. existing: [{id,name,price,category}]. Mavjud mahsulotning kategoriyasi Excel'dagiga
    yangilanadi (kategoriya nomi mavjud kategoriya bilan solishtiriladi: ЕZOZ = EZOZ)."""
    by_name = {norm_name(r["name"]): r for r in existing}
    cats = {cat_key(r["category"]): r["category"] for r in existing}
    updates, inserts, missing, moved, new_cats = [], [], [], [], []
    updated = unchanged = 0
    for name, price, cat in rows:
        target = None
        if cat:
            k = cat_key(cat)
            if k not in cats:
                cats[k] = cat
                new_cats.append(cat)
            target = cats[k]
        r = by_name.get(norm_name(name))
        if r:
            new_price = price if (price and price != r["price"]) else None
            new_cat = target if (target and target != r["category"]) else None
            if new_price is not None:
                updated += 1
            elif new_cat is None:
                unchanged += 1
            if new_price is not None or new_cat is not None:
                updates.append((r["id"], new_price, new_cat))
            if new_cat is not None:
                moved.append({"name": r["name"], "from": r["category"], "to": new_cat})
        elif price:  # narxsiz va bazada yo'q mahsulotlar e'tiborga olinmaydi
            if add_new:
                inserts.append((name, price, target or "Прочие"))
            else:
                missing.append(name)
    return {"updates": updates, "inserts": inserts, "missing": missing, "moved": moved,
            "new_cats": new_cats, "updated": updated, "unchanged": unchanged}


async def h_import_excel(request):
    raw, add_new = None, False
    reader = await request.multipart()
    async for part in reader:
        if part.name == "file":
            raw = await part.read(decode=False)
        elif part.name == "add_new":
            add_new = (await part.text()).strip() in ("1", "true", "yes")
    if not raw:
        return err("no_file")
    try:
        rows = parse_price_excel(raw)
    except Exception as e:
        log.warning("Excel o'qilmadi: %s", e)
        return err("bad_excel")
    if not rows:
        return err("empty_excel")

    async with db.pool.acquire() as c:
        existing = [dict(r) for r in await c.fetch("select id,name,price,category from products")]
        plan = plan_import(rows, existing, add_new)
        next_id = (await c.fetchval("select coalesce(max(id),0) from products")) + 1
        next_sort = (await c.fetchval("select coalesce(max(sort),0) from products")) + 1
        async with c.transaction():
            for pid, price, cat in plan["updates"]:
                await c.execute(
                    "update products set price=coalesce($1,price), category=coalesce($2,category) where id=$3",
                    price, cat, pid)
            for name, price, cat in plan["inserts"]:
                await c.execute(
                    "insert into products(id,name,price,category,sort) values($1,$2,$3,$4,$5)",
                    next_id, name, price, cat, next_sort)
                next_id += 1
                next_sort += 1
    return web.json_response({
        "ok": True, "updated": plan["updated"], "unchanged": plan["unchanged"], "added": len(plan["inserts"]),
        "not_found": plan["missing"][:50], "not_found_count": len(plan["missing"]),
        "recategorized": len(plan["moved"]), "moved": plan["moved"][:40], "new_categories": plan["new_cats"],
    })


async def h_product_image(request):
    pid = int(request.match_info["id"])
    ctype = request.content_type or ""
    ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}.get(ctype)
    if not ext:
        return err("bad_image_type")
    body = await request.read()
    if not body or len(body) > 8 * 1024 * 1024:
        return err("bad_image_size")
    if not (SUPABASE_URL and SUPABASE_KEY):
        return err("storage_not_configured", 500)

    path = f"p{pid}_{int(time.time())}.{ext}"
    url = f"{SUPABASE_URL}/storage/v1/object/{SUPABASE_BUCKET}/{path}"
    headers = {"apikey": SUPABASE_KEY, "Content-Type": ctype, "x-upsert": "true"}
    if SUPABASE_KEY.startswith("eyJ"):  # eski JWT kalit; yangi sb_secret_ kalit faqat apikey'da
        headers["Authorization"] = f"Bearer {SUPABASE_KEY}"
    try:
        async with aiohttp.ClientSession() as s:
            async with s.post(url, data=body, headers=headers) as r:
                if r.status >= 300:
                    text = (await r.text())[:300]
                    log.error("Storage xatosi %s: %s", r.status, text)
                    return web.json_response(
                        {"ok": False, "error": "upload_failed", "detail": f"{r.status}: {text}"}, status=502)
    except Exception as e:
        log.error("Storage ulanish xatosi: %s", e)
        return web.json_response({"ok": False, "error": "upload_failed", "detail": str(e)[:200]}, status=502)
    public = f"{SUPABASE_URL}/storage/v1/object/public/{SUPABASE_BUCKET}/{path}"
    async with db.pool.acquire() as c:
        r = await c.fetchrow("update products set image_url=$1 where id=$2 returning id", public, pid)
    return web.json_response({"ok": True, "image": public}) if r else err("not_found", 404)


# ------------------------------------------------------------------ ADMIN: sozlamalar, buyurtmalar
async def h_settings_get(request):
    return web.json_response({"ok": True, **(await order_status())})


async def h_settings_set(request):
    data = await request.json()
    async with db.pool.acquire() as c:
        for key in ("open_time", "cutoff_time"):
            if key in data:
                v = str(data[key])
                if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", v):
                    return err("bad_time")
                await c.execute(
                    "insert into settings(key,value) values($1,$2) "
                    "on conflict (key) do update set value=excluded.value", key, v,
                )
        if "order_days" in data:  # zakas qabul qilinadigan kunlar (admin belgilaydi)
            raw = data["order_days"]
            if not isinstance(raw, list):
                return err("bad_input")
            try:
                days = sorted({datetime.strptime(str(x), "%Y-%m-%d").date().isoformat() for x in raw})
            except ValueError:
                return err("bad_date")
            await c.execute(
                "insert into settings(key,value) values('order_days',$1) "
                "on conflict (key) do update set value=excluded.value", json.dumps(days),
            )
    return web.json_response({"ok": True, **(await order_status())})


async def pending_warnings(c, rows) -> dict:
    """Tasdiq kutayotgan zakaslar uchun ogohlantirishlar {order_id: {avto: matn}}: Shinaray/Isuzi shu kunga boshqa
    filialga tasdiqlangan bo'lsa yoki Shinaray shu haftada shu filialga tasdiqlangan bo'lsa. Admin baribir tanlay oladi."""
    pend = [r for r in rows if r["status"] == "pending"]
    if not pend:
        return {}
    dates = [r["order_date"] for r in pend]
    conf = await c.fetch(
        "select o.user_id,o.order_date,o.vehicle,u.name from orders o join users u on u.id=o.user_id "
        "where o.status='confirmed' and o.vehicle in ('shinaray','isuzi') and o.order_date between $1 and $2",
        min(dates) - timedelta(days=6), max(dates) + timedelta(days=6))
    out = {}
    for r in pend:
        ws = r["order_date"] - timedelta(days=r["order_date"].weekday())
        w: dict[str, list] = {}
        for k in conf:
            if k["order_date"] == r["order_date"] and k["user_id"] != r["user_id"]:
                w.setdefault(k["vehicle"], []).append(f"shu kunga {k['name']} filialiga tasdiqlangan")
            if (k["vehicle"] == "shinaray" and k["user_id"] == r["user_id"]
                    and ws <= k["order_date"] < ws + timedelta(days=7)):
                w.setdefault("shinaray", []).append("bu filialga shu haftada Shinaray tasdiqlangan (haftada 1 marta)")
        if w:
            out[r["id"]] = {veh: "; ".join(msgs) for veh, msgs in w.items()}
    return out


async def h_admin_orders(request):
    date = request.query.get("date")
    status = request.query.get("status")
    where, args = [], []
    if date:
        try:
            args.append(datetime.strptime(date, "%Y-%m-%d").date())  # asyncpg date kutadi, str emas
        except ValueError:
            return err("bad_date")
        where.append(f"o.order_date=${len(args)}")
    if status in ("pending", "confirmed"):
        args.append(status)
        where.append(f"o.status=${len(args)}")
    q = ("select o.*, u.name as uname, u.phone as uphone from orders o join users u on u.id=o.user_id "
         + ("where " + " and ".join(where) + " " if where else "")
         + ("order by o.order_date, o.id " if status == "pending" else "order by o.id desc ") + "limit 200")
    async with db.pool.acquire() as c:
        rows = await c.fetch(q, *args)
        warn = await pending_warnings(c, rows)
    out = []
    for r in rows:
        x = order_json(r)
        x["user"] = f"{r['uname']} (+998{r['uphone']})"
        x["warn"] = warn.get(r["id"], {})
        out.append(x)
    return web.json_response({"ok": True, "orders": out})


async def h_order_confirm(request):
    """Admin zakasga avto tanlab tasdiqlaydi — shundan keyingina zakas guruhga boradi."""
    oid = int(request.match_info["id"])
    data = await request.json()
    veh = data.get("vehicle")
    if veh not in VEHICLES:
        return err("bad_vehicle")
    try:
        n = int(data.get("vehicle_n") or 1)
    except (TypeError, ValueError):
        return err("bad_vehicle")
    async with db.pool.acquire() as c:
        r = await c.fetchrow(
            "select o.*, u.name as uname, u.phone as uphone from orders o join users u on u.id=o.user_id "
            "where o.id=$1", oid)
        if not r:
            return err("not_found", 404)
        if r["status"] != "pending":
            return err("already_confirmed", 409)
        items = json.loads(r["items"]) if isinstance(r["items"], str) else r["items"]
        eq = eq_kg(items)
        if not fits(veh, n, eq):
            return err("vehicle_mismatch")
        upd = await c.fetchval(
            "update orders set vehicle=$1, vehicle_n=$2, status='confirmed', confirmed_at=now() "
            "where id=$3 and status='pending' returning id", veh, n, oid)
        if not upd:
            return err("already_confirmed", 409)
    V = VEHICLES[veh]
    sent = await notify(order_text(
        oid, r["order_date"], V["name"] + (f" ×{n}" if n > 1 else ""), LABO_CAP * n if veh == "labo" else V["max_kg"],
        {"name": r["uname"], "phone": r["uphone"]}, items, float(r["total_kg"]), float(r["total_dona"]), eq))
    return web.json_response({"ok": True, "sent": sent})


async def h_order_del(request):
    oid = int(request.match_info["id"])
    async with db.pool.acquire() as c:
        async with c.transaction():
            r = await c.fetchval("delete from orders where id=$1 returning id", oid)
            if r:  # zakas qo'shimcha huquq bilan berilgan bo'lsa, huquq qaytariladi
                await c.execute("update grants set used_order_id=null where used_order_id=$1", oid)
    return web.json_response({"ok": True}) if r else err("not_found", 404)


async def h_control(request):
    """Kontrol: filial × avto jadvali (kg) — hamma foydalanuvchiga. So'rovlar/huquqlar — faqat adminga."""
    date = request.query.get("date") or datetime.now(TIMEZONE).date().isoformat()
    try:
        d = datetime.strptime(date, "%Y-%m-%d").date()
    except ValueError:
        return err("bad_date")
    is_admin = bool(request["user"]["is_admin"])
    async with db.pool.acquire() as c:
        orders = await c.fetch(
            "select o.id,o.user_id,o.vehicle,o.vehicle_n,o.items,u.name,u.phone from orders o "
            "join users u on u.id=o.user_id where o.order_date=$1 and o.status='confirmed' order by o.id", d)
        if is_admin:
            users = await c.fetch("select id,name,phone from users where active order by name")
            grants = await c.fetch(
                "select g.id,g.user_id,g.vehicle,u.name from grants g join users u on u.id=g.user_id "
                "where g.used_order_id is null order by g.id desc")
            reqs = await c.fetch(
                "select r.id,r.user_id,r.text,r.created_at,u.name from requests r "
                "join users u on u.id=r.user_id order by r.id desc limit 30")
    rows: dict[int, dict] = {}
    totals = {k: 0.0 for k in VEHICLES}
    for o in orders:
        v = o["vehicle"]
        if v not in VEHICLES:
            continue  # eski (avtosiz) buyurtmalar «Buyurtmalar» bo'limida ko'rinadi
        items = json.loads(o["items"]) if isinstance(o["items"], str) else o["items"]
        kg = eq_kg(items)
        row = rows.setdefault(o["user_id"], {
            "user_id": o["user_id"], "name": o["name"],
            **{k: {"kg": 0.0, "n": 0, "cars": 0} for k in VEHICLES},
        })
        row[v]["kg"] = round(row[v]["kg"] + kg, 2)
        row[v]["n"] += 1
        row[v]["cars"] += o["vehicle_n"] or 1
        totals[v] = round(totals[v] + kg, 2)
    totals["all"] = round(sum(totals.values()), 2)
    out = {"ok": True, "date": d.isoformat(), "rows": list(rows.values()), "totals": totals}
    if is_admin:
        out.update({
            "users": [{"id": x["id"], "name": x["name"], "phone": "+998" + x["phone"]} for x in users],
            "grants": [{"id": g["id"], "user_id": g["user_id"], "name": g["name"], "vehicle": g["vehicle"]}
                       for g in grants],
            "requests": [{"id": q["id"], "user_id": q["user_id"], "name": q["name"], "text": q["text"],
                          "created_at": q["created_at"].astimezone(TIMEZONE).strftime("%d.%m %H:%M")}
                         for q in reqs],
        })
    return web.json_response(out)


async def h_grant_add(request):
    data = await request.json()
    veh = data.get("vehicle")
    try:
        uid = int(data.get("user_id"))
        cnt = max(1, min(int(data.get("count") or 1), 5))
    except (TypeError, ValueError):
        return err("bad_input")
    if veh not in VEHICLES:
        return err("bad_vehicle")
    async with db.pool.acquire() as c:
        if not await c.fetchval("select 1 from users where id=$1", uid):
            return err("not_found", 404)
        for _ in range(cnt):
            await c.execute("insert into grants(user_id,vehicle) values($1,$2)", uid, veh)
    return web.json_response({"ok": True})


async def h_grant_del(request):
    gid = int(request.match_info["id"])
    async with db.pool.acquire() as c:
        r = await c.fetchval("delete from grants where id=$1 and used_order_id is null returning id", gid)
    return web.json_response({"ok": True}) if r else err("not_found", 404)


async def h_health(request):
    try:
        await db.ping()
        return web.json_response({"ok": True})
    except Exception as e:
        return web.json_response({"ok": False, "error": str(e)}, status=500)


APP_DIR = "../app/www"


async def h_app_index(request):
    return web.FileResponse(f"{APP_DIR}/index.html", headers={"Cache-Control": "no-cache"})


def add_routes(app: web.Application):
    r = app.router
    r.add_get("/app", h_app_index)
    r.add_get("/app/", h_app_index)
    r.add_static("/app/", path=APP_DIR, show_index=False)
    r.add_post("/api/v2/tg-login", h_tg_login)
    r.add_get("/health", h_health)
    r.add_post("/api/v2/check", h_check)
    r.add_post("/api/v2/login", h_login)
    r.add_get("/api/v2/me", h_me)
    r.add_get("/api/v2/products", h_products)
    r.add_post("/api/v2/order", h_order)
    r.add_get("/api/v2/orders", h_my_orders)
    r.add_get("/api/v2/admin/users", h_users)
    r.add_post("/api/v2/admin/users", h_user_add)
    r.add_post("/api/v2/admin/users/{id}", h_user_update)
    r.add_get("/api/v2/admin/products", h_admin_products)
    r.add_post("/api/v2/admin/products/{id}", h_product_update)
    r.add_post("/api/v2/admin/products/{id}/image", h_product_image)
    r.add_post("/api/v2/admin/import-excel", h_import_excel)
    r.add_get("/api/v2/admin/settings", h_settings_get)
    r.add_post("/api/v2/admin/settings", h_settings_set)
    r.add_get("/api/v2/admin/orders", h_admin_orders)
    r.add_get("/api/v2/vehicles", h_vehicles)
    r.add_post("/api/v2/request", h_request)
    r.add_get("/api/v2/control", h_control)
    r.add_post("/api/v2/admin/orders/{id}/confirm", h_order_confirm)
    r.add_delete("/api/v2/admin/orders/{id}", h_order_del)
    r.add_post("/api/v2/admin/grants", h_grant_add)
    r.add_delete("/api/v2/admin/grants/{id}", h_grant_del)
