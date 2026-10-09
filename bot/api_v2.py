"""Android ilova uchun yangi API (/api/v2/...). Eski Telegram Mini App yo'llari (/api/products, /api/order) o'zgarmaydi."""
import hashlib
import hmac
import io
import json
import logging
import re
import secrets
import time
from datetime import datetime, timedelta

import aiohttp
from aiohttp import web

import db
from verify import verify_init_data
from config import (
    SECRET_KEY, TIMEZONE, ADMIN_CHAT_ID, ADMIN_PHONE,
    SUPABASE_URL, SUPABASE_KEY, SUPABASE_BUCKET,
)

log = logging.getLogger("zakas-api")

BOT = None  # main.py o'rnatadi
TOKEN_TTL = 30 * 24 * 3600
FAILS: dict[str, list[float]] = {}  # telefon -> xato urinishlar vaqti


# ------------------------------------------------------------------ yordamchilar
def make_token(uid: int) -> str:
    exp = int(time.time()) + TOKEN_TTL
    msg = f"{uid}.{exp}"
    sig = hmac.new(SECRET_KEY.encode(), msg.encode(), hashlib.sha256).hexdigest()
    return f"{msg}.{sig}"


def parse_token(token: str) -> int | None:
    try:
        uid, exp, sig = token.split(".")
        msg = f"{uid}.{exp}"
        good = hmac.new(SECRET_KEY.encode(), msg.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(good, sig) or int(exp) < time.time():
            return None
        return int(uid)
    except Exception:
        return None


def err(code: str, status: int = 400) -> web.Response:
    return web.json_response({"ok": False, "error": code}, status=status)


def user_json(u, with_code: bool = False) -> dict:
    d = {
        "id": u["id"], "name": u["name"], "phone": "+998" + u["phone"],
        "is_admin": u["is_admin"], "active": u["active"],
    }
    if with_code:
        d["code"] = u["code"]
    return d


def tg_user_id(init_data: str) -> int | None:
    """Telegram initData to'g'ri va yangi bo'lsa — foydalanuvchi ID'si, aks holda None."""
    pairs = verify_init_data(init_data or "")
    if pairs is None:
        return None
    try:
        if time.time() - int(pairs.get("auth_date", 0)) > 7 * 24 * 3600:
            return None
        return int(json.loads(pairs.get("user", "{}")).get("id", 0)) or None
    except (ValueError, TypeError):
        return None


def hhmm_to_min(s: str) -> int:
    h, m = s.split(":")
    return int(h) * 60 + int(m)


async def order_status() -> dict:
    s = await db.get_settings()
    now = datetime.now(TIMEZONE)
    cur = now.hour * 60 + now.minute
    is_open = hhmm_to_min(s["open_time"]) <= cur < hhmm_to_min(s["cutoff_time"])
    return {"open_time": s["open_time"], "cutoff_time": s["cutoff_time"], "order_open": is_open}


def norm_name(s: str) -> str:
    return re.sub(r"[\W_]+", "", str(s).lower())


# ------------------------------------------------------------------ middleware
CORS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Headers": "Authorization, Content-Type",
    "Access-Control-Allow-Methods": "GET, POST, PUT, DELETE, OPTIONS",
}


@web.middleware
async def cors_mw(request, handler):
    if not request.path.startswith("/api/v2"):
        return await handler(request)
    if request.method == "OPTIONS":
        return web.Response(status=204, headers=CORS)
    resp = await handler(request)
    resp.headers.update(CORS)
    return resp


PUBLIC = {"/api/v2/check", "/api/v2/login", "/api/v2/tg-login"}


@web.middleware
async def auth_mw(request, handler):
    path = request.path
    if not path.startswith("/api/v2") or path in PUBLIC:
        return await handler(request)
    auth = request.headers.get("Authorization", "")
    uid = parse_token(auth.removeprefix("Bearer ").strip())
    if uid is None:
        return err("unauthorized", 401)
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where id=$1", uid)
    if not u or not u["active"]:
        return err("not_allowed", 403)
    if path.startswith("/api/v2/admin") and not u["is_admin"]:
        return err("forbidden", 403)
    request["user"] = u
    return await handler(request)


# ------------------------------------------------------------------ kirish
async def h_check(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where phone=$1", phone)
    if not u or not u["active"]:
        return web.json_response({"ok": True, "status": "denied"})
    return web.json_response({"ok": True, "status": "new" if not u["code"] else "has_code"})


async def h_login(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    code = str(data.get("code") or "").strip()

    now = time.time()
    FAILS[phone] = [t for t in FAILS.get(phone, []) if now - t < 600]
    if len(FAILS[phone]) >= 5:
        return err("too_many_attempts", 429)

    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where phone=$1", phone)
        if not u or not u["active"]:
            return err("not_allowed", 403)
        if not u["code"]:  # birinchi kirish: foydalanuvchi o'z kodini yozadi
            if len(code) < 4:
                return err("code_too_short", 400)
            await c.execute("update users set code=$1 where id=$2", code, u["id"])
            u = await c.fetchrow("select * from users where id=$1", u["id"])
        elif not hmac.compare_digest(u["code"], code):
            FAILS[phone].append(now)
            return err("wrong_code", 401)
        tg = tg_user_id(str(data.get("initData") or ""))
        if tg:  # Telegram ichidan kirdi: keyingi safar kod so'ramasligi uchun bog'laymiz
            await c.execute("update users set tg_id=null where tg_id=$1", tg)
            await c.execute("update users set tg_id=$1 where id=$2", tg, u["id"])
    FAILS.pop(phone, None)
    return web.json_response({"ok": True, "token": make_token(u["id"]), "user": user_json(u)})


async def h_tg_login(request):
    """Telegram Mini App: bog'langan foydalanuvchi (yoki ADMIN_CHAT_ID egasi) kodsiz kiradi."""
    data = await request.json()
    tg_id = tg_user_id(str(data.get("initData") or ""))
    if tg_id is None:
        return err("auth_failed", 401)
    async with db.pool.acquire() as c:
        u = await c.fetchrow("select * from users where tg_id=$1 and active", tg_id)
        if not u and ADMIN_CHAT_ID and tg_id == ADMIN_CHAT_ID:
            u = await c.fetchrow(
                "select * from users where phone=$1 and is_admin and active", db.norm_phone(ADMIN_PHONE))
    if not u:
        return err("not_allowed", 403)
    return web.json_response({"ok": True, "token": make_token(u["id"]), "user": user_json(u)})


async def h_me(request):
    return web.json_response({"ok": True, "user": user_json(request["user"])})


# ------------------------------------------------------------------ katalog va buyurtma
async def h_products(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch(
            "select id,name,price,category,image_url from products where active order by sort,id"
        )
    catalog: dict[str, list] = {}
    for r in rows:
        catalog.setdefault(r["category"], []).append(
            {"id": r["id"], "name": r["name"], "price": r["price"], "image": r["image_url"]}
        )
    return web.json_response({"ok": True, "catalog": catalog, **(await order_status())})


async def h_order(request):
    u = request["user"]
    data = await request.json()
    st = await order_status()
    if not st["order_open"]:
        return err("closed", 403)

    today = datetime.now(TIMEZONE).date()
    try:
        d = datetime.strptime(str(data.get("date") or today.isoformat()), "%Y-%m-%d").date()
    except ValueError:
        return err("bad_date")
    if d < today or d > today + timedelta(days=14):
        return err("bad_date")

    async with db.pool.acquire() as c:
        prods = {r["id"]: r for r in await c.fetch("select id,name,price from products where active")}
    lines, items, kg, dona = [], [], 0.0, 0.0
    for it in data.get("items") or []:
        p = prods.get(it.get("id"))
        try:
            qty = float(it.get("qty", 0))
        except (TypeError, ValueError):
            continue
        if not p or qty <= 0:
            continue
        unit = "dona" if it.get("unit") == "dona" else "kg"
        kg += qty if unit == "kg" else 0
        dona += qty if unit == "dona" else 0
        items.append({"id": p["id"], "name": p["name"], "price": p["price"], "qty": qty, "unit": unit})
        lines.append(f"• {p['name']} — {qty:g} {unit} ({p['price']:,} so'm/{unit})".replace(",", " "))
    if not items:
        return err("empty_order")

    async with db.pool.acquire() as c:
        oid = await c.fetchval(
            "insert into orders(user_id,order_date,items,total_kg,total_dona) "
            "values($1,$2,$3::jsonb,$4,$5) returning id",
            u["id"], d, json.dumps(items, ensure_ascii=False), kg, dona,
        )

    if ADMIN_CHAT_ID and BOT:
        try:
            text = (
                f"🆕 <b>Yangi buyurtma #{oid}</b> (ilova)\n"
                f"📅 Sana: {d.strftime('%d.%m.%Y')}\n"
                f"👤 {u['name']} (+998{u['phone']})\n\n"
                f"<b>Mahsulotlar:</b>\n" + "\n".join(lines) + "\n\n"
                f"<b>Jami:</b> {kg:g} kg, {dona:g} dona"
            )
            await BOT.send_message(ADMIN_CHAT_ID, text, parse_mode="HTML")
        except Exception as e:
            log.warning("Telegramga yuborib bo'lmadi: %s", e)
    return web.json_response({"ok": True, "order_id": oid})


async def h_my_orders(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch(
            "select id,order_date,created_at,items,total_kg,total_dona from orders "
            "where user_id=$1 order by id desc limit 50", request["user"]["id"],
        )
    return web.json_response({"ok": True, "orders": [order_json(r) for r in rows]})


def order_json(r) -> dict:
    return {
        "id": r["id"], "date": r["order_date"].isoformat(),
        "created_at": r["created_at"].astimezone(TIMEZONE).strftime("%Y-%m-%d %H:%M"),
        "items": json.loads(r["items"]) if isinstance(r["items"], str) else r["items"],
        "total_kg": float(r["total_kg"]), "total_dona": float(r["total_dona"]),
    }


# ------------------------------------------------------------------ ADMIN: foydalanuvchilar
async def h_users(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch("select * from users order by id")
    return web.json_response({"ok": True, "users": [user_json(r, True) for r in rows]})


async def h_user_add(request):
    data = await request.json()
    phone = db.norm_phone(data.get("phone"))
    name = str(data.get("name") or "").strip()
    if len(phone) != 9 or not name:
        return err("bad_input")
    async with db.pool.acquire() as c:
        if await c.fetchval("select 1 from users where phone=$1", phone):
            return err("exists", 409)
        uid = await c.fetchval(
            "insert into users(phone,name) values($1,$2) returning id", phone, name
        )
    return web.json_response({"ok": True, "id": uid})


async def h_user_update(request):
    uid = int(request.match_info["id"])
    data = await request.json()
    sets, vals = [], []

    def add(col, val):
        vals.append(val)
        sets.append(f"{col}=${len(vals)}")

    if "name" in data:
        add("name", str(data["name"]).strip())
    if "active" in data:
        if uid == request["user"]["id"] and not data["active"]:
            return err("cannot_disable_self")
        add("active", bool(data["active"]))
    if "is_admin" in data:
        if uid == request["user"]["id"] and not data["is_admin"]:
            return err("cannot_demote_self")
        add("is_admin", bool(data["is_admin"]))
    if data.get("generate_code") or "code" in data:
        add("tg_id", None)  # kod o'zgarsa Telegram'dan qayta kod bilan kirishi kerak
    if data.get("generate_code"):
        add("code", f"{secrets.randbelow(10000):04d}")
    elif "code" in data:  # None yoki "" => kodni tozalash (foydalanuvchi o'zi qayta yozadi)
        c_ = str(data["code"] or "").strip()
        if c_ and len(c_) < 4:
            return err("code_too_short")
        add("code", c_ or None)
    if not sets:
        return err("nothing_to_update")
    vals.append(uid)
    async with db.pool.acquire() as c:
        u = await c.fetchrow(
            f"update users set {', '.join(sets)} where id=${len(vals)} returning *", *vals
        )
    if not u:
        return err("not_found", 404)
    return web.json_response({"ok": True, "user": user_json(u, True)})


# ------------------------------------------------------------------ ADMIN: mahsulot, narx
async def h_admin_products(request):
    async with db.pool.acquire() as c:
        rows = await c.fetch("select * from products order by sort,id")
    return web.json_response({"ok": True, "products": [
        {"id": r["id"], "name": r["name"], "price": r["price"], "category": r["category"],
         "image": r["image_url"], "active": r["active"]} for r in rows]})


async def h_product_update(request):
    pid = int(request.match_info["id"])
    data = await request.json()
    sets, vals = [], []
    for col, cast in (("price", int), ("name", str), ("category", str), ("active", bool)):
        if col in data:
            vals.append(cast(data[col]))
            sets.append(f"{col}=${len(vals)}")
    if not sets:
        return err("nothing_to_update")
    vals.append(pid)
    async with db.pool.acquire() as c:
        r = await c.fetchrow(
            f"update products set {', '.join(sets)} where id=${len(vals)} returning id", *vals
        )
    return web.json_response({"ok": True}) if r else err("not_found", 404)


def parse_price_excel(raw: bytes) -> list[tuple[str, int]]:
    """Excel'dan (nom, narx) juftliklarini oladi. Sarlavha qatorini o'zi topadi."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    name_keys = ("наимен", "номи", "nom", "name", "товар", "mahsulot", "product")
    price_keys = ("цена", "прайс", "narx", "price", "нарх", "стоим")
    name_col = price_col = hdr = None
    for i, row in enumerate(rows[:15]):
        n = p = None
        for j, cell in enumerate(row):
            t = str(cell or "").lower()
            if n is None and any(k in t for k in name_keys):
                n = j
            if p is None and any(k in t for k in price_keys):
                p = j
        if n is not None and p is not None:
            name_col, price_col, hdr = n, p, i
            break
    if hdr is None:  # zaxira: birinchi matn ustuni va birinchi son ustuni
        hdr = 0
        sample = rows[1] if len(rows) > 1 else []
        name_col = next((j for j, v in enumerate(sample) if isinstance(v, str)), 0)
        price_col = next((j for j, v in enumerate(sample) if isinstance(v, (int, float))), 1)
    out = []
    for row in rows[hdr + 1:]:
        if len(row) <= max(name_col, price_col):
            continue
        name, price = row[name_col], row[price_col]
        if not name:
            continue
        try:
            price = int(round(float(str(price).replace(" ", "").replace(",", "."))))
        except (TypeError, ValueError):
            continue
        if price > 0:
            out.append((str(name).strip(), price))
    return out


async def h_import_excel(request):
    raw, add_new = None, False
    reader = await request.multipart()
    async for part in reader:
        if part.name == "file":
            raw = await part.read(decode=False)
        elif part.name == "add_new":
            add_new = (await part.text()).strip() in ("1", "true", "yes")
    if not raw:
        return err("no_file")
    try:
        pairs = parse_price_excel(raw)
    except Exception as e:
        log.warning("Excel o'qilmadi: %s", e)
        return err("bad_excel")
    if not pairs:
        return err("empty_excel")

    updated, unchanged, added, missing = 0, 0, 0, []
    async with db.pool.acquire() as c:
        existing = {norm_name(r["name"]): r for r in await c.fetch("select id,name,price from products")}
        next_id = (await c.fetchval("select coalesce(max(id),0) from products")) + 1
        next_sort = (await c.fetchval("select coalesce(max(sort),0) from products")) + 1
        async with c.transaction():
            for name, price in pairs:
                r = existing.get(norm_name(name))
                if r:
                    if r["price"] != price:
                        await c.execute("update products set price=$1 where id=$2", price, r["id"])
                        updated += 1
                    else:
                        unchanged += 1
                elif add_new:
                    await c.execute(
                        "insert into products(id,name,price,category,sort) values($1,$2,$3,'Прочие',$4)",
                        next_id, name, price, next_sort,
                    )
                    next_id += 1
                    next_sort += 1
                    added += 1
                else:
                    missing.append(name)
    return web.json_response({
        "ok": True, "updated": updated, "unchanged": unchanged, "added": added,
        "not_found": missing[:50], "not_found_count": len(missing),
    })


async def h_product_image(request):
    pid = int(request.match_info["id"])
    ctype = request.content_type or ""
    ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}.get(ctype)
    if not ext:
        return err("bad_image_type")
    body = await request.read()
    if not body or len(body) > 8 * 1024 * 1024:
        return err("bad_image_size")
    if not (SUPABASE_URL and SUPABASE_KEY):
        return err("storage_not_configured", 500)

    path = f"p{pid}_{int(time.time())}.{ext}"
    url = f"{SUPABASE_URL}/storage/v1/object/{SUPABASE_BUCKET}/{path}"
    headers = {"apikey": SUPABASE_KEY, "Content-Type": ctype, "x-upsert": "true"}
    if SUPABASE_KEY.startswith("eyJ"):  # eski JWT kalit; yangi sb_secret_ kalit faqat apikey'da
        headers["Authorization"] = f"Bearer {SUPABASE_KEY}"
    try:
        async with aiohttp.ClientSession() as s:
            async with s.post(url, data=body, headers=headers) as r:
                if r.status >= 300:
                    text = (await r.text())[:300]
                    log.error("Storage xatosi %s: %s", r.status, text)
                    return web.json_response(
                        {"ok": False, "error": "upload_failed", "detail": f"{r.status}: {text}"}, status=502)
    except Exception as e:
        log.error("Storage ulanish xatosi: %s", e)
        return web.json_response({"ok": False, "error": "upload_failed", "detail": str(e)[:200]}, status=502)
    public = f"{SUPABASE_URL}/storage/v1/object/public/{SUPABASE_BUCKET}/{path}"
    async with db.pool.acquire() as c:
        r = await c.fetchrow("update products set image_url=$1 where id=$2 returning id", public, pid)
    return web.json_response({"ok": True, "image": public}) if r else err("not_found", 404)


# ------------------------------------------------------------------ ADMIN: sozlamalar, buyurtmalar
async def h_settings_get(request):
    return web.json_response({"ok": True, **(await order_status())})


async def h_settings_set(request):
    data = await request.json()
    async with db.pool.acquire() as c:
        for key in ("open_time", "cutoff_time"):
            if key in data:
                v = str(data[key])
                if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", v):
                    return err("bad_time")
                await c.execute(
                    "insert into settings(key,value) values($1,$2) "
                    "on conflict (key) do update set value=excluded.value", key, v,
                )
    return web.json_response({"ok": True, **(await order_status())})


async def h_admin_orders(request):
    date = request.query.get("date")
    q = ("select o.*, u.name as uname, u.phone as uphone from orders o join users u on u.id=o.user_id "
         + ("where o.order_date=$1::date " if date else "") + "order by o.id desc limit 200")
    async with db.pool.acquire() as c:
        rows = await c.fetch(q, *( [date] if date else [] ))
    out = []
    for r in rows:
        d = order_json(r)
        d["user"] = f"{r['uname']} (+998{r['uphone']})"
        out.append(d)
    return web.json_response({"ok": True, "orders": out})


async def h_health(request):
    try:
        await db.ping()
        return web.json_response({"ok": True})
    except Exception as e:
        return web.json_response({"ok": False, "error": str(e)}, status=500)


APP_DIR = "../app/www"


async def h_app_index(request):
    return web.FileResponse(f"{APP_DIR}/index.html", headers={"Cache-Control": "no-cache"})


def add_routes(app: web.Application):
    r = app.router
    r.add_get("/app", h_app_index)
    r.add_get("/app/", h_app_index)
    r.add_static("/app/", path=APP_DIR, show_index=False)
    r.add_post("/api/v2/tg-login", h_tg_login)
    r.add_get("/health", h_health)
    r.add_post("/api/v2/check", h_check)
    r.add_post("/api/v2/login", h_login)
    r.add_get("/api/v2/me", h_me)
    r.add_get("/api/v2/products", h_products)
    r.add_post("/api/v2/order", h_order)
    r.add_get("/api/v2/orders", h_my_orders)
    r.add_get("/api/v2/admin/users", h_users)
    r.add_post("/api/v2/admin/users", h_user_add)
    r.add_post("/api/v2/admin/users/{id}", h_user_update)
    r.add_get("/api/v2/admin/products", h_admin_products)
    r.add_post("/api/v2/admin/products/{id}", h_product_update)
    r.add_post("/api/v2/admin/products/{id}/image", h_product_image)
    r.add_post("/api/v2/admin/import-excel", h_import_excel)
    r.add_get("/api/v2/admin/settings", h_settings_get)
    r.add_post("/api/v2/admin/settings", h_settings_set)
    r.add_get("/api/v2/admin/orders", h_admin_orders)
